from datetime import datetime, timedelta, timezone
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from pytest import mark

from tests.test import ClockifyTestCase
from clockify.session import ClockifySession
from clockify.model.reports_model import SummaryReport, SummaryReportGetParams, SummaryFilter
from clockify.model.time_entry_model import (
    CreateTimeEntryDTO,
    TimeEntry,
    TimeEntryGetParams,
)

class TestReports(ClockifyTestCase):
    _TEST_TIME_ENTRY_NAME = 'Test Time Entry'
    
    @classmethod
    def setUpClass(cls) -> None:
        cls.session = ClockifySession(cls.KEY)
        
        utc_timezone = timezone(timedelta(hours=0), 'UTC')
        start = datetime.now(utc_timezone)
        end = start + timedelta(minutes=30)
        
        time_entry_data = CreateTimeEntryDTO(
            start=start, end=end, description=cls._TEST_TIME_ENTRY_NAME
        )
        
        cls.session.time_entry.add_time_entry(
            cls.WORKSPACE, cls.session.get_current_user().id_, time_entry_data
        )
        
        return super().setUpClass()

    @classmethod
    def tearDownClass(cls) -> None:
        return super().tearDownClass()

    @mark.filterwarnings("ignore:Workbook contains no default style")
    def test_summary_xlsx_report(self):
        date_range_start = datetime.now().strftime('%Y-%m-%dT00:00:00Z')
        date_range_end = datetime.now().strftime('%Y-%m-%dT23:59:59Z')
        summary_report_get_params = SummaryReportGetParams(
            summary_filter=SummaryFilter(
                groups=['USER', 'TIMEENTRY'],
                sort_column='GROUP',
                summary_chart_type='PROJECT'
                ),
            export_type='XLSX',
            date_range_start=date_range_start,
            date_range_end=date_range_end
            )
        binart_excel_data = self.session.reports.summary(self.WORKSPACE, summary_report_get_params)
        self.assertIsInstance(binart_excel_data, bytes)
        wb = load_workbook(BytesIO(binart_excel_data))
        COLUMN_TIME_ENTRY_DESCRIPTION_B = 'B';
        COLUMN_DECIMAL_TIME_D = 'D';
        self.assertEqual(wb.worksheets[0].cell(row=3, column=column_index_from_string(COLUMN_TIME_ENTRY_DESCRIPTION_B)).value, self._TEST_TIME_ENTRY_NAME)
        self.assertEqual(wb.worksheets[0].cell(row=3, column=column_index_from_string(COLUMN_DECIMAL_TIME_D)).value, 0.5)
        wb.close()